from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
REF = ROOT / 'docs' / 'Kweneng_Public_Website_Quotation_2026-08-06.docx'
OUT = ROOT / 'docs' / 'Swaneng_Website_Development_Quotation_2026-08-10.docx'
copy2(REF, OUT)
doc = Document(OUT)

# Replace the template body while retaining its section, styles, theme, and footer.
body = doc._element.body
sect_pr = body.sectPr
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

BLUE = '2E74B5'
NAVY = '17365D'
PALE = 'EAF2F8'
LIGHT = 'F4F7FA'
TEXT = RGBColor(32, 38, 49)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def add_heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.add_run(text)
    keep_with_next(p)
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def add_label_value(label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r=p.add_run(label + ': '); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    p.add_run(value)

# Branded running header.
header = doc.sections[0].header
hp = header.paragraphs[0]
hp.clear(); hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run('DIGITALCORE BUSINESS SYSTEMS  |  WEBSITE SOLUTIONS')
hr.bold=True; hr.font.size=Pt(9); hr.font.color.rgb=RGBColor.from_string(BLUE)

# Cover/title block.
p=doc.add_paragraph(style='Title')
p.add_run('Website Development\nQuotation')
p.paragraph_format.space_after=Pt(6)
p=doc.add_paragraph()
r=p.add_run('A modern, professional and mobile-friendly school website')
r.font.size=Pt(13); r.font.color.rgb=RGBColor(92,102,115)
p.paragraph_format.space_after=Pt(18)

meta=doc.add_table(rows=5, cols=2)
meta.alignment=WD_TABLE_ALIGNMENT.LEFT; meta.autofit=False
labels=['Prepared for','Prepared by','Date','Quotation validity','Service']
vals=['Swaneng English Medium Primary School','DigitalCore Business Systems','10 August 2026','30 days','School website development']
for i,(lab,val) in enumerate(zip(labels,vals)):
    meta.columns[0].width=Inches(1.65); meta.columns[1].width=Inches(4.75)
    for c in meta.rows[i].cells: set_cell_margins(c,100,120,100,120); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(meta.cell(i,0), BLUE if i==0 else PALE)
    set_cell_shading(meta.cell(i,1), LIGHT if i%2 else 'FFFFFF')
    p1=meta.cell(i,0).paragraphs[0]; p1.paragraph_format.space_after=Pt(0)
    rr=p1.add_run(lab.upper()); rr.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor.from_string(NAVY)
    p2=meta.cell(i,1).paragraphs[0]; p2.paragraph_format.space_after=Pt(0)
    rr=p2.add_run(val); rr.bold=(i==0); rr.font.size=Pt(10.5); rr.font.color.rgb=TEXT

doc.add_paragraph()
add_heading('1. School website development',1)
p=doc.add_paragraph('DigitalCore Business Systems is pleased to provide this quotation for the development of a modern, professional and mobile-friendly website for Swaneng English Medium Primary School.')
p.paragraph_format.space_after=Pt(7)
p=doc.add_paragraph('The website will present clear information about the school and will be easy for authorised school staff to manage and update.')
p.paragraph_format.space_after=Pt(8)

add_heading('Included features',2)
for item in [
    'Professional website design', 'Mobile-friendly and responsive layout', 'Home page',
    'About the School page', 'Academic information pages', 'Contact page',
    'News and announcements section', 'Photo gallery', 'Social media integration',
    'Contact form', 'Administrator login for website management']:
    add_bullet(item)

add_heading('2. Website content management',1)
doc.add_paragraph('The website will include a user-friendly administration panel that allows school staff to update the website without technical knowledge.')
add_heading('Authorised staff will be able to',2)
for item in ['Update website content and text','Upload and manage photos','Publish news and announcements','Update contact information','Add and edit pages as required']:
    add_bullet(item)

add_heading('3. Quotation',1)
qt=doc.add_table(rows=2, cols=2)
qt.alignment=WD_TABLE_ALIGNMENT.LEFT; qt.autofit=False
widths=[Inches(4.9), Inches(1.5)]
for row in qt.rows:
    for i,c in enumerate(row.cells): c.width=widths[i]; set_cell_margins(c,140,160,140,160); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for i,text in enumerate(['Description','Amount (BWP)']):
    set_cell_shading(qt.cell(0,i),BLUE); p=qt.cell(0,i).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run(text); r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
qt.cell(1,0).paragraphs[0].add_run('School Website Development').bold=True
p=qt.cell(1,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('P7,500'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(NAVY)

offer=doc.add_table(rows=1,cols=1); offer.alignment=WD_TABLE_ALIGNMENT.LEFT
cell=offer.cell(0,0); set_cell_shading(cell,PALE); set_cell_margins(cell,180,180,180,180)
p=cell.paragraphs[0]; r=p.add_run('SPECIAL PACKAGE OFFER'); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
p=cell.add_paragraph('When the website is purchased together with the School Management System, the website price is ')
r=p.add_run('P5,500'); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(NAVY)
p.add_run('.')

add_heading('4. Delivery period',1)
p=doc.add_paragraph('Estimated delivery time is '); r=p.add_run('2 to 4 weeks'); r.bold=True; p.add_run(' from the date of approval and receipt of all required content.')

add_heading('5. Payment terms',1)
add_bullet('50% deposit before commencement of work.')
add_bullet('50% upon completion and handover.')

add_heading('6. Notes',1)
for item in [
    'Domain registration and hosting fees are excluded unless otherwise agreed.',
    'Any additional functionality requested outside the agreed scope may be quoted separately.',
    'Basic training on how to update the website will be provided.'
]: add_bullet(item)

add_heading('7. Acceptance',1)
sig=doc.add_table(rows=4,cols=2); sig.alignment=WD_TABLE_ALIGNMENT.LEFT; sig.autofit=False
heads=['For Swaneng English Medium Primary School','For DigitalCore Business Systems']
for i,h in enumerate(heads):
    sig.columns[i].width=Inches(3.2); set_cell_shading(sig.cell(0,i),PALE); set_cell_margins(sig.cell(0,i),140,140,140,140)
    r=sig.cell(0,i).paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
for row,label in enumerate(['Name: __________________________','Signature: ______________________','Date: ___________________________'], start=1):
    for c in sig.rows[row].cells:
        set_cell_margins(c,80,140,80,140); c.paragraphs[0].paragraph_format.space_after=Pt(0); c.paragraphs[0].add_run(label)

# Keep footer page-number field from the source; normalize body font.
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10); normal.font.color.rgb=TEXT
normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1.04
for style_name in ('Heading 1','Heading 2'):
    st=doc.styles[style_name]
    st.paragraph_format.space_before=Pt(8 if style_name=='Heading 1' else 5)
    st.paragraph_format.space_after=Pt(3)
doc.core_properties.title='Website Development Quotation - Swaneng English Medium Primary School'
doc.core_properties.subject='School website development quotation'
doc.core_properties.author='DigitalCore Business Systems'
doc.core_properties.comments='Prepared 10 August 2026'
doc.save(OUT)
print(OUT)
