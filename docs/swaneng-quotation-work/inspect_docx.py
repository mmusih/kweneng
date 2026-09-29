from docx import Document
from pathlib import Path
import json, sys

path = Path(sys.argv[1])
doc = Document(path)
data = {
    "sections": [{
        "width": s.page_width, "height": s.page_height,
        "top": s.top_margin, "bottom": s.bottom_margin,
        "left": s.left_margin, "right": s.right_margin,
        "header": s.header_distance, "footer": s.footer_distance,
    } for s in doc.sections],
    "paragraphs": [], "tables": [],
    "headers": [], "footers": []
}
for i,p in enumerate(doc.paragraphs):
    data["paragraphs"].append({"i":i,"text":p.text,"style":p.style.name,"align":str(p.alignment),"runs":[{"text":r.text,"bold":r.bold,"italic":r.italic,"size":r.font.size.pt if r.font.size else None,"font":r.font.name,"color":str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None} for r in p.runs]})
for ti,t in enumerate(doc.tables):
    rows=[]
    for ri,row in enumerate(t.rows):
        rows.append([{"text":c.text,"paras":[{"text":p.text,"style":p.style.name,"align":str(p.alignment)} for p in c.paragraphs]} for c in row.cells])
    data["tables"].append({"i":ti,"style":t.style.name if t.style else None,"rows":rows})
for s in doc.sections:
    data["headers"].append([p.text for p in s.header.paragraphs])
    data["footers"].append([p.text for p in s.footer.paragraphs])
print(json.dumps(data, indent=2))
